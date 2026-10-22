#!/usr/bin/env python3
"""
FORCE — Gera planilha .xlsx pronta para abrir no Google Sheets.
Execute: python gerar_planilha.py
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from datetime import date

# ---------- Paleta ----------
AZUL_ESCURO = "12395C"
AZUL_MEDIO  = "266BA6"
AZUL_CLARO  = "D6E7F7"
CINZA_CLARO = "F4F4F4"
AMARELO     = "FFE9B0"
CINZA_TXT   = "595959"
BRANCO      = "FFFFFF"

# ---------- Estilos ----------
def fill(hex_):  return PatternFill("solid", fgColor=hex_)
def fonte(tam=10, bold=False, cor="000000", italic=False):
    return Font(name="Calibri", size=tam, bold=bold, color=cor, italic=italic)
def alin(h="left", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

BORDA_FINA = Border(*[Side(style="thin", color="D9D9D9")]*4)


def titulo(ws, rng, texto, cor_fundo=AZUL_ESCURO, cor_txt=BRANCO, tam=14):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = texto
    c.fill = fill(cor_fundo)
    c.font = fonte(tam, bold=True, cor=cor_txt)
    c.alignment = alin("left", "center")


def faixa(ws, rng, texto, cor_fundo=AZUL_CLARO, cor_txt=AZUL_ESCURO, tam=10, italic=True):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    c.value = texto
    c.fill = fill(cor_fundo)
    c.font = fonte(tam, italic=italic, cor=cor_txt)
    c.alignment = alin("left", "center")


def cabecalho(ws, linha, colunas):
    for i, nome in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=i, value=nome)
        c.fill = fill(AZUL_MEDIO)
        c.font = fonte(9, bold=True, cor=BRANCO)
        c.alignment = alin("center", "center")
        c.border = BORDA_FINA


def linha_ajuda(ws, linha, descricoes):
    for i, txt in enumerate(descricoes, start=1):
        c = ws.cell(row=linha, column=i, value=txt)
        c.fill = fill(CINZA_CLARO)
        c.font = fonte(8, italic=True, cor=CINZA_TXT)
        c.alignment = alin("left", "top")
        c.border = BORDA_FINA


def ajustar_larguras(ws, larguras):
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def configurar_impressao(ws, ultima_col, ultima_linha=200):
    """Paisagem + ajusta a 1 página de largura ao imprimir."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_area = f"A1:{ultima_col}{ultima_linha}"


# ============================================================
# ABA GUIA
# ============================================================
def criar_guia(wb):
    ws = wb.active
    ws.title = "Guia"

    for col, w in zip("ABCD", [22, 38, 38, 16]):
        ws.column_dimensions[col].width = w

    titulo(ws, "A1:D1", "🏢   FORCE — SISTEMA DE GESTÃO DE ATIVOS", tam=16)
    ws.row_dimensions[1].height = 40

    titulo(ws, "A2:D2",
           "Planilha-mãe do controle patrimonial, sub-ativos, eventos e indicadores",
           cor_fundo=AZUL_MEDIO, tam=10)
    ws["A2"].font = fonte(10, italic=True, cor=BRANCO)
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 6

    texto = (
        "👋  Bem-vindo(a)!\n"
        "Esta planilha é a base do sistema de gestão de ativos da FORCE. "
        "Use-a para controlar veículos, implementos, pneus, baterias, "
        "geradores embarcados e qualquer equipamento com rastreabilidade individual. "
        "Antes de começar, leia as descrições das abas e as Regras de Ouro."
    )
    faixa(ws, "A4:D4", texto, tam=10, italic=False)
    ws["A4"].alignment = alin("left", "center")
    ws.row_dimensions[4].height = 60
    ws.row_dimensions[5].height = 6

    titulo(ws, "A6:D6", "🗂️   ABAS DESTA PLANILHA", tam=12)
    ws.row_dimensions[6].height = 24

    for i, nome in enumerate(["Aba", "O que é", "Quando usar", "Responsável"], start=1):
        c = ws.cell(row=7, column=i, value=nome)
        c.fill = fill(AZUL_MEDIO)
        c.font = fonte(10, bold=True, cor=BRANCO)
        c.alignment = alin("center", "center")
    ws.row_dimensions[7].height = 22

    abas = [
        ("📖  Guia", "Esta aba. Explica o sistema e as regras.",
         "Sempre que tiver dúvida sobre onde registrar algo.", "Todos"),
        ("🚛  Ativos", "Cadastro dos bens principais: caminhões, carretas, vans, motos, geradores.",
         "Ao adquirir, transferir ou baixar um ativo.", "Gestão / Frota"),
        ("🔧  Sub-ativos", "Componentes vinculados: pneus, baterias, eletrônicos.",
         "Ao instalar, remover ou substituir um componente rastreável.", "Manutenção"),
        ("📋  Eventos", "Manutenções, vistorias, avarias, movimentações.",
         "A cada ocorrência.", "Manutenção / Gestão"),
        ("📚  Listas", "Opções padronizadas usadas nas listas suspensas.",
         "Ao adicionar nova classe, base, oficina ou tipo.", "Admin"),
    ]
    r = 8
    for i, linha in enumerate(abas):
        for j, txt in enumerate(linha, start=1):
            c = ws.cell(row=r, column=j, value=txt)
            c.fill = fill(CINZA_CLARO if i % 2 == 0 else BRANCO)
            c.alignment = alin("left", "center")
            c.font = fonte(9)
            c.border = BORDA_FINA
        ws.cell(row=r, column=1).font = fonte(10, bold=True, cor=AZUL_ESCURO)
        ws.row_dimensions[r].height = 34
        r += 1

    r += 1

    titulo(ws, f"A{r}:D{r}", "🔑   REGRAS DE OURO", tam=12)
    ws.row_dimensions[r].height = 24
    r += 1

    regras = [
        "Uma linha por registro. Nunca junte dois ativos na mesma linha.",
        'Nunca deixe campo obrigatório em branco. Use "N/A" ou "A definir".',
        "Use as listas suspensas para padronizar.",
        "Asset ID e Sub-ativo ID são únicos. Nunca repita.",
        "Toda movimentação gera uma linha na aba Eventos.",
        "Atualize o Cadastro de Ativos sempre que houver evento.",
        "Mantenha os links de evidência atualizados.",
    ]
    for i, regra in enumerate(regras, start=1):
        a = ws.cell(row=r, column=1, value=str(i))
        a.fill = fill(AZUL_CLARO); a.font = fonte(10, bold=True, cor=AZUL_ESCURO)
        a.alignment = alin("center", "center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(row=r, column=2, value=regra)
        c.fill = fill(CINZA_CLARO if i % 2 else BRANCO)
        c.font = fonte(9); c.alignment = alin("left", "center")
        ws.row_dimensions[r].height = 20
        r += 1

    r += 1
    titulo(ws, f"A{r}:D{r}", "🗺️   FLUXO DE USO", tam=12)
    ws.row_dimensions[r].height = 24
    r += 1
    faixa(ws, f"A{r}:D{r}",
          "1. Cadastrar Ativo  ➜  2. Vincular Sub-ativos  ➜  3. Registrar Eventos  ➜  4. Atualizar Ativo",
          cor_fundo=AMARELO, cor_txt="4D3300", tam=11, italic=False)
    ws.row_dimensions[r].height = 26
    r += 2

    rodape = f"FORCE — Sistema de Gestão de Ativos  •  v1.0  •  {date.today().strftime('%d/%m/%Y')}"
    titulo(ws, f"A{r}:D{r}", rodape, tam=8)
    ws[f"A{r}"].font = fonte(8, italic=True, cor=BRANCO)
    ws.row_dimensions[r].height = 18

    configurar_impressao(ws, "D", r)


# ============================================================
# ABAS DE DADOS
# ============================================================
def criar_aba(wb, nome, icone, titulo_txt, descricao, colunas, descricoes, larguras):
    ws = wb.create_sheet(nome)
    n = len(colunas)
    ult = get_column_letter(n)

    titulo(ws, f"A1:{ult}1", f"{icone}   {titulo_txt}", tam=13)
    ws.row_dimensions[1].height = 30

    faixa(ws, f"A2:{ult}2", descricao)
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 5

    cabecalho(ws, 4, colunas)
    ws.row_dimensions[4].height = 28

    linha_ajuda(ws, 5, descricoes)
    ws.row_dimensions[5].height = 32

    ajustar_larguras(ws, larguras)
    ws.freeze_panes = "A6"
    configurar_impressao(ws, ult)
    return ws


def criar_ativos(wb):
    colunas = ["Asset ID","Placa","Chassi","Tipo de Ativo","Marca","Modelo","Ano",
               "Próprio/Locado","Locador","Contrato","Base","Responsável Operacional",
               "Km/Horímetro Atual","Data Última Atualização","Classe","Status",
               "Data Última Inspeção","Data Próxima Inspeção","Pendências Abertas",
               "Link Passaporte Técnico","Observações"]
    descricoes = ["ID único interno.","Placa do veículo.","Número do chassi.",
                  "Cavalo, carreta, van, moto, gerador...","Fabricante.","Modelo.",
                  "Ano de fabricação.","Próprio ou locado.","Nome do locador.",
                  "Contrato vinculado.","Base atual.","Quem opera.",
                  "Km/horímetro atual.","Última atualização.","A/B/C/D.",
                  "Situação atual.","Última inspeção.","Próxima inspeção.",
                  "Resumo curto.","Link do passaporte.","Campo livre."]
    larguras = [12, 10, 18, 14, 12, 14, 7, 12, 12, 12, 8, 16,
                13, 12, 7, 13, 12, 12, 22, 16, 22]
    ws = criar_aba(wb, "Ativos", "🚛", "CADASTRO DE ATIVOS",
                   "Um ativo por linha. Bens principais: caminhões, carretas, vans, motos, geradores.",
                   colunas, descricoes, larguras)
    add_validacoes(ws, {
        "D": "=Listas!$A$5:$A$30",
        "H": "=Listas!$H$5:$H$10",
        "I": "=Listas!$E$5:$E$30",
        "K": "=Listas!$F$5:$F$30",
        "O": "=Listas!$B$5:$B$10",
        "P": "=Listas!$C$5:$C$20",
    })


def criar_sub_ativos(wb):
    colunas = ["Sub-ativo ID","Tipo de Sub-ativo","Ativo Pai (Asset ID)","Posição",
               "Número de Série","Marca","Modelo","DOT","Data Instalação",
               "Km Instalação","Data Retirada","Km Retirada","Motivo Retirada",
               "Status","Observações"]
    descricoes = ["ID único (ex: PNEU-00123).","Pneu, bateria, eletrônico...",
                  "Asset ID do ativo pai.","Posição física.",
                  "Número de série.","Fabricante.","Modelo.",
                  "Semana/ano (pneus).","Data instalação.","Km na instalação.",
                  "Data retirada.","Km na retirada.","Motivo.","Situação.","Campo livre."]
    larguras = [14, 16, 16, 16, 14, 12, 12, 7, 12, 12, 12, 12, 18, 13, 20]
    ws = criar_aba(wb, "Sub-ativos", "🔧", "CADASTRO DE SUB-ATIVOS",
                   "Um sub-ativo por linha. Componentes rastreáveis vinculados a um ativo principal.",
                   colunas, descricoes, larguras)
    add_validacoes(ws, {
        "B": "=Listas!$J$5:$J$30",
        "N": "=Listas!$I$5:$I$20",
    })


def criar_eventos(wb):
    colunas = ["Evento ID","Data","Ativo (Asset ID)","Sub-ativo ID","Tipo de Evento",
               "Descrição","Diagnóstico Técnico","Classe no Momento","Recomendação",
               "Prazo","Responsável Provável","OS SGLOC","Status","Validado por",
               "Evidência (link)","Observações"]
    descricoes = ["ID único (EV-0001).","Data do evento.","Asset ID.",
                  "Sub-ativo envolvido.","Preventiva, corretiva, vistoria...",
                  "O que aconteceu.","Diagnóstico.","Classe no momento.",
                  "Ação recomendada.","Prazo.","Responsável.","Nº OS SGLOC.",
                  "Status do evento.","Quem validou.","Link de evidência.","Campo livre."]
    larguras = [11, 10, 14, 12, 16, 24, 24, 12, 20, 10, 14, 11, 14, 12, 16, 20]
    ws = criar_aba(wb, "Eventos", "📋", "MOVIMENTAÇÕES / EVENTOS",
                   "Uma linha por evento. Manutenções, vistorias, avarias e movimentações.",
                   colunas, descricoes, larguras)
    add_validacoes(ws, {
        "E": "=Listas!$D$5:$D$30",
        "H": "=Listas!$B$5:$B$10",
        "M": "=Listas!$C$5:$C$20",
    })


def add_validacoes(ws, mapa):
    for col, ref in mapa.items():
        dv = DataValidation(type="list", formula1=ref, allow_blank=True, showDropDown=False)
        dv.error = "Escolha uma opção da lista."
        dv.errorTitle = "Valor inválido"
        ws.add_data_validation(dv)
        dv.add(f"{col}6:{col}1000")


# ============================================================
# ABA LISTAS
# ============================================================
def criar_listas(wb):
    ws = wb.create_sheet("Listas")

    headers = ["Tipo de Ativo","Classe","Status","Tipo de Evento","Locador",
               "Base","Oficina","Próprio/Locado","Status Sub-ativo","Tipo de Sub-ativo"]
    listas = {
        "Tipo de Ativo":     ["Caminhão","Carreta","Van","Moto","Gerador","Climatização","Nobreak","Implemento","Outro"],
        "Classe":            ["A","B","C","D"],
        "Status":            ["Operando","Parado","Em manutenção","Aguardando validação","Bloqueado","Baixado"],
        "Tipo de Evento":    ["Manutenção preventiva","Manutenção corretiva","Vistoria","Avaria","Movimentação de pneu","Troca de bateria","Instalação de equipamento","Outro"],
        "Locador":           ["Localiza","Unidas","Movida","Outro"],
        "Base":              ["SP","RJ","MG","PR","SC","BA","PE","Outro"],
        "Oficina":           ["Oficina Zé","ENDY","Outro"],
        "Próprio/Locado":    ["Próprio","Locado"],
        "Status Sub-ativo":  ["Em uso","Estoque","Condenado","Recapagem","Baixado"],
        "Tipo de Sub-ativo": ["Pneu","Bateria","Gerador embarcado","Eletrônico","Outro"],
    }

    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.fill = fill(AZUL_MEDIO)
        c.font = fonte(9, bold=True, cor=BRANCO)
        c.alignment = alin("center", "center")
        c.border = BORDA_FINA

    for i, h in enumerate(headers, start=1):
        for j, val in enumerate(listas[h]):
            c = ws.cell(row=5+j, column=i, value=val)
            c.fill = fill(CINZA_CLARO)
            c.font = fonte(9)
            c.alignment = alin("left", "center")

    for i in range(1, len(headers)+1):
        ws.column_dimensions[get_column_letter(i)].width = 18

    titulo(ws, "A1:J1", "📚   LISTAS AUXILIARES", tam=13)
    ws.row_dimensions[1].height = 30
    faixa(ws, "A2:J2",
          "Opções padronizadas usadas nas listas suspensas das outras abas. "
          "Adicione novas opções abaixo de cada lista — elas aparecem automaticamente.")
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 5
    ws.freeze_panes = "A5"
    configurar_impressao(ws, "J")


# ============================================================
# MAIN
# ============================================================
def main():
    wb = Workbook()
    criar_guia(wb)
    criar_listas(wb)
    criar_ativos(wb)
    criar_sub_ativos(wb)
    criar_eventos(wb)

    ordem = ["Guia", "Ativos", "Sub-ativos", "Eventos", "Listas"]
    wb._sheets = [wb[n] for n in ordem]

    nome_arquivo = "FORCE_Ativos.xlsx"
    wb.save(nome_arquivo)
    print(f"✅ Arquivo criado: {nome_arquivo}")
    print("👉 Suba no Drive e abra com Google Sheets.")


if __name__ == "__main__":
    main()